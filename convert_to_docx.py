# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def sanitize(text):
    return "".join(ch for ch in text if ch in ('\t', '\n', '\r') or 0x20 <= ord(ch) <= 0xD7FF or 0xE000 <= ord(ch) <= 0xFFFD)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def flush_table(lines_to_flush, doc):
    if not lines_to_flush:
        return
    rows_data = []
    for l in lines_to_flush:
        l_str = l.strip()
        if l_str.startswith('|') and l_str.endswith('|'):
            parts = [sanitize(p.strip()) for p in l_str[1:-1].split('|')]
            if all(set(p).issubset({'-', ':', ' '}) for p in parts):
                continue
            rows_data.append(parts)
    if not rows_data:
        return
    num_cols = max(len(r) for r in rows_data)
    tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = 'Table Grid'
    for r_idx, row in enumerate(rows_data):
        for c_idx, cell_text in enumerate(row):
            if c_idx < num_cols:
                cell = tbl.cell(r_idx, c_idx)
                cell.text = cell_text.replace('**', '').replace('*', '')
                set_cell_margins(cell, 80, 80, 120, 120)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.space_after = Pt(2)
                    for run in p.runs:
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(10)
                        if r_idx == 0:
                            run.bold = True
    doc.add_paragraph()

def flush_code(c_lines, doc):
    if not c_lines:
        return
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(sanitize('\n'.join(c_lines)))
    run.font.name = 'Courier New'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(30, 30, 30)

def create_report_docx():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(6)

    with open("TY_Project_Report.md", "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    for line in lines:
        raw_line = line.rstrip('\r\n')
        stripped = raw_line.strip()

        if '<div style="page-break-after: always;"></div>' in stripped:
            if in_table:
                flush_table(table_lines, doc)
                table_lines = []
                in_table = False
            doc.add_page_break()
            continue

        if stripped.startswith('```'):
            if in_code_block:
                flush_code(code_lines, doc)
                code_lines = []
                in_code_block = False
            else:
                if in_table:
                    flush_table(table_lines, doc)
                    table_lines = []
                    in_table = False
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(raw_line)
            continue

        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            continue
        else:
            if in_table:
                flush_table(table_lines, doc)
                table_lines = []
                in_table = False

        if not stripped:
            continue

        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(sanitize(stripped[2:]))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(18)
            run.bold = True
            if 'PROJECT REPORT' in stripped or 'EDUVAULT' in stripped:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(sanitize(stripped[3:]))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(15)
            run.bold = True
            if any(k in stripped for k in ['CERTIFICATE', 'DECLARATION', 'ACKNOWLEDGEMENT', 'ABSTRACT']):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(sanitize(stripped[4:]))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(13)
            run.bold = True
        elif stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(sanitize(stripped[5:]))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.bold = True
        elif stripped.startswith('---'):
            continue
        else:
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(6)
            clean_text = sanitize(stripped.replace('**', '').replace('*', ''))
            run = p.add_run(clean_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)

    if in_table:
        flush_table(table_lines, doc)
    if in_code_block:
        flush_code(code_lines, doc)

    output_path = "TY_Project_Report.docx"
    doc.save(output_path)
    print(f"Successfully created {output_path}")

if __name__ == '__main__':
    create_report_docx()
